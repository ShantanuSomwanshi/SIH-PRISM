"""
Turn an uploaded tender document into searchable text.

The problem statement asks for "product descriptions, technical
specifications, or tender documents" as input. Typing into a box covers
the first two; this covers the third.

PDFs go through the same extractor the corpus uses, so a scanned tender
is OCR'd automatically - exactly as a scanned standard is.

An upload is read entirely in memory. Nothing from it is written to disk -
no temp file and no OCR cache entry - so a tender does not outlive the
request that carried it.
"""

import io
from pathlib import Path

from backend.config import MAX_DOC_PAGES, MAX_UPLOAD_BYTES

SUPPORTED_SUFFIXES = {".pdf", ".docx", ".txt", ".md"}


class UnsupportedDocument(Exception):
    """Raised when the file cannot be read. The message is shown to the user."""


def _read_txt(data: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _read_docx(data: bytes) -> str:
    try:
        import docx  # python-docx
    except ImportError as exc:
        raise UnsupportedDocument(
            "Reading .docx files needs the python-docx package. "
            "Install it with: pip install python-docx"
        ) from exc

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]

    # Tender requirements live in tables at least as often as in paragraphs.
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    return "\n".join(parts)


def _read_pdf(data: bytes) -> tuple:
    """
    Read an uploaded PDF from memory.

    This used to write the bytes to a temp file called "upload.pdf" and
    pass that to extract_pdf(). The OCR cache is named after the file, so
    every scanned tender's text was saved to ocr_cache/upload.json and left
    there - overwritten by the next upload, and mixed up between two
    uploads arriving together. extract_pdf_bytes() never caches.
    """
    from backend.pdf_extract import extract_pdf_bytes

    result = extract_pdf_bytes(data, max_pages=MAX_DOC_PAGES)

    text = "\n\n".join(page.text for page in result.pages if page.text)
    return text, {
        "pages_read": len(result.pages),
        "total_pages": result.total_pages,
        "ocr_pages": result.ocr_pages,
    }


def extract_upload(filename: str, data: bytes) -> tuple:
    """Read an uploaded file and return (text, info)."""
    if not data:
        raise UnsupportedDocument("The uploaded file is empty.")
    if len(data) > MAX_UPLOAD_BYTES:
        raise UnsupportedDocument(
            f"File is too large. The limit is "
            f"{MAX_UPLOAD_BYTES // (1024 * 1024)} MB."
        )

    suffix = Path(filename or "").suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise UnsupportedDocument(
            f"Cannot read '{suffix or 'that file type'}'. Supported types: "
            + ", ".join(sorted(SUPPORTED_SUFFIXES))
        )

    info = {"filename": filename, "type": suffix}

    if suffix == ".pdf":
        text, pdf_info = _read_pdf(data)
        info.update(pdf_info)
    elif suffix == ".docx":
        text = _read_docx(data)
    else:
        text = _read_txt(data)

    text = text.strip()
    if not text:
        raise UnsupportedDocument(
            "No readable text was found in that file. If it is a scan, the "
            "pages may be too poor for OCR."
        )

    info["characters"] = len(text)
    return text, info
